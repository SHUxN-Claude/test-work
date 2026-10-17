# -*- coding: utf-8 -*-
"""比較ブック(前回今回比較.xlsm)の「抽出_要対応」箇所について、
前後のキロ程を含めた前回・今回の重ね合わせチャートを作成する。

使い方:
  python build_charts.py 比較ブック.xlsm 出力先フォルダ [--window 100] [--max-shift 5]

位置補正なし版・位置補正あり版の2ファイルを出力する。
"""
import argparse
import math
import os
import re

import openpyxl
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.title import Title
from openpyxl.chart.text import Text
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import (CharacterProperties, Paragraph,
                                   ParagraphProperties, RegularTextRun)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# チャートに並べる項目(上から順に)。システムのチャートと同じ並び。
ITEMS = ["高低（左）", "高低（右）", "通り（左）", "通り（右）", "平面性", "水準", "軌間"]
# 位置ズレ推定に使う項目(波形の形がはっきりしている項目)
SHIFT_ITEMS = ["高低（左）", "高低（右）", "通り（左）", "通り（右）"]
# 「位置ズレの可能性あり」とする目安(参考判定):ズレ量0.5m以上 かつ 一致度が15%以上改善
SHIFT_MIN_M = 0.5
SHIFT_MIN_IMPROVE = 0.15

KM_FMT = '0"k"000'          # 84870 → 84k870
COLOR_CURR = "0000FF"       # 今回:青太線
COLOR_PREV = "FF0000"       # 前回:赤細線
COLOR_SHIFT = "00A040"      # 前回(位置補正後):緑破線
COLOR_LOW = "FFA000"        # 抽出 10mm以上15mm未満:オレンジ縦線
COLOR_HIGH = "FF0000"       # 抽出 15mm以上:赤縦線

RULE_TEXT = ("〔変位進み量が15㎜以上〕\n検査後15日以内に調査\n"
             "　→変位進みを確認した場合、30日以内に処置\n\n"
             "〔変位進み量が10㎜以上15㎜未満〕\n検査後30日以内に調査\n"
             "　→変位進みを確認した場合、60日以内に処置")

EMU_PT = 12700
PANEL_ROWS = [6, 6, 6, 6, 6, 6, 8]   # 行2〜45(44行)を7段に配分。最下段は目盛分だけ高く
DATA_COL = 24                         # X列からチャート用データを置く(印刷範囲A:Vの外)


# ---------------------------------------------------------------- 読み込み
def fnum(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.startswith('="') and s.endswith('"'):
        s = s[2:-1]
    try:
        return float(s)
    except ValueError:
        return None


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["比較"].iter_rows(values_only=True))
    ext = list(wb["抽出_要対応"].iter_rows(values_only=True))
    wb.close()
    h = {k: i for i, k in enumerate(rows[0])}
    eh = {k: i for i, k in enumerate(ext[0])}

    def track_of(r, hh):
        return (r[hh["線名"]], r[hh["線路名"]], r[hh["本準側名"]])

    tracks = {}
    for r in rows[1:]:
        if r[h["キロ程"]] is None:
            continue
        t = track_of(r, h)
        km = int(round(fnum(r[h["キロ程"]])))
        d = tracks.setdefault(t, {})
        d[km] = {"駅名称": r[h.get("駅名称", h["線名"])] if "駅名称" in h else "",
                 **{f"{it}前回": fnum(r[h[f"{it}前回"]]) for it in ITEMS},
                 **{f"{it}今回": fnum(r[h[f"{it}今回"]]) for it in ITEMS}}
    points = []
    for n, r in enumerate(ext[1:], 1):
        if r[eh["キロ程"]] is None:
            continue
        hits = [x for x in str(r[eh["該当項目"]] or "").split("、")
                if x and not x.startswith("(")]
        adv = {}
        for it in hits:
            dv = fnum(r[eh[f"{it}差分"]])
            if dv is not None:
                adv[it] = dv
        points.append({
            "番号": n,
            "No": r[eh["No"]],
            "track": track_of(r, eh),
            "駅名称": r[eh["駅名称"]] if "駅名称" in eh else "",
            "km": int(round(fnum(r[eh["キロ程"]]))),
            "判定": r[eh["判定"]],
            "該当項目": hits,
            "差分": adv,
            "前回実績番号": r[eh["前回実績番号"]],
            "今回実績番号": r[eh["今回実績番号"]],
        })
    return tracks, points


def make_groups(points, window):
    """同じ線路で前後window mの範囲が重なる抽出箇所を1グループにまとめる。"""
    groups = []
    for p in points:
        g = groups[-1] if groups else None
        if g and g["track"] == p["track"] and p["km"] - window <= g["points"][-1]["km"] + window:
            g["points"].append(p)
        else:
            groups.append({"track": p["track"], "points": [p]})
    return groups


def sheet_name(nums):
    if len(nums) == 1:
        return str(nums[0])
    if nums == list(range(nums[0], nums[-1] + 1)):
        return f"{nums[0]}～{nums[-1]}"
    return "・".join(str(n) for n in nums)[:31]


# ---------------------------------------------------------------- 位置ズレ推定
def interp(series, x):
    """1m間隔データ series(dict km->値) の x における線形補間値。"""
    k0 = math.floor(x)
    t = x - k0
    a, b = series.get(k0), series.get(k0 + 1)
    if t < 1e-9:
        return a
    if a is None or b is None:
        return None
    return a + (b - a) * t


def estimate_shift(data, points, max_shift, core=30):
    """前回を d m ずらしたとき(前回(k-d) と 今回(k) を比較)に最も一致する d を探す。
    抽出箇所の前後core mを評価範囲とし、SHIFT_ITEMSの差の二乗和で評価する。"""
    ks = sorted({k for p in points for k in range(p["km"] - core, p["km"] + core + 1)})
    prev = {it: {k: v[f"{it}前回"] for k, v in data.items() if v[f"{it}前回"] is not None}
            for it in SHIFT_ITEMS}
    curr = {it: {k: v[f"{it}今回"] for k, v in data.items() if v[f"{it}今回"] is not None}
            for it in SHIFT_ITEMS}

    def score(d):
        ss, n = 0.0, 0
        for it in SHIFT_ITEMS:
            for k in ks:
                c = curr[it].get(k)
                pv = interp(prev[it], k - d)
                if c is not None and pv is not None:
                    ss += (c - pv) ** 2
                    n += 1
        return math.sqrt(ss / n) if n else float("inf")

    steps = int(round(max_shift * 10))
    cands = [i / 10 for i in range(-steps, steps + 1)]
    best = min(cands, key=lambda d: (round(score(d), 6), abs(d)))
    return best, score(0.0), score(best)


def corrected_adv(data, p, d):
    """位置補正後の進み量(今回(k) − 前回(k−d))。該当項目ごと。"""
    out = {}
    k = p["km"]
    for it in p["該当項目"]:
        c = data.get(k, {}).get(f"{it}今回")
        prev = {kk: v[f"{it}前回"] for kk, v in data.items() if v[f"{it}前回"] is not None}
        pv = interp(prev, k - d)
        if c is not None and pv is not None:
            out[it] = round(c - pv, 1)
    return out


def judge(v):
    a = abs(v)
    return "15mm以上" if a >= 15 else ("10mm以上15mm未満" if a >= 10 else "10mm未満")


# ---------------------------------------------------------------- チャート部品
def rich(size=800, bold=False, color=None):
    cp = CharacterProperties(sz=size, b=bold)
    if color:
        cp.solidFill = color
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)
    para.r = []
    return RichText(p=[para])


def title(text, size=900, bold=True, color=None):
    cp = CharacterProperties(sz=size, b=bold)
    if color:
        cp.solidFill = color
    paras = [Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=line)])
             for line in text.split("\n")]
    t = Title(tx=Text(rich=RichText(p=paras)))
    return t


def line_series(ws, xcol, ycol, r0, r1, name, color, width_pt, dash=None):
    xs = Reference(ws, min_col=xcol, min_row=r0, max_row=r1)
    ys = Reference(ws, min_col=ycol, min_row=r0, max_row=r1)
    s = Series(ys, xs, title=name)
    s.marker.symbol = "none"
    s.smooth = False
    s.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=color, w=int(width_pt * EMU_PT), prstDash=dash))
    return s


def anchor(c0, r0, c1, r1):
    a = TwoCellAnchor(editAs="twoCell")
    a._from = AnchorMarker(col=c0, row=r0)
    a.to = AnchorMarker(col=c1, row=r1)
    return a


# ---------------------------------------------------------------- シート作成
def setup_sheet(ws):
    ws.sheet_format.defaultRowHeight = 13.5
    ws.row_dimensions[1].height = 119.25
    for c in range(1, 23):
        ws.column_dimensions[get_column_letter(c)].width = 9.0
    ws.merge_cells("A1:B1")
    ws.merge_cells("C1:P1")
    ws.merge_cells("Q1:V1")
    a1 = ws["A1"]
    a1.value = '=MID(CELL("filename",A1),FIND("]",CELL("filename",A1))+1,31)'
    a1.font = Font(name="ＭＳ Ｐゴシック", size=24, bold=True, color="FF0000")
    a1.alignment = Alignment(horizontal="center", vertical="center")
    c1 = ws["C1"]
    c1.font = Font(name="ＭＳ Ｐゴシック", size=24, bold=True)
    c1.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    q1 = ws["Q1"]
    q1.value = RULE_TEXT
    q1.font = Font(name="ＭＳ Ｐゴシック", size=12)
    q1.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.print_area = "A1:V45"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = ws.page_margins.right = 0.315
    ws.page_margins.top = ws.page_margins.bottom = 0.354
    ws.sheet_view.view = "pageBreakPreview"
    ws.sheet_view.zoomScaleSheetLayoutView = 100


def write_group_sheet(wb, g, data, window, with_shift, shift_info):
    nums = [p["番号"] for p in g["points"]]
    ws = wb.create_sheet(sheet_name(nums))
    setup_sheet(ws)
    kmin_all, kmax_all = min(data), max(data)
    lo = max(g["points"][0]["km"] - window, kmin_all)
    hi = min(g["points"][-1]["km"] + window, kmax_all)
    d = shift_info["d"] if with_shift else 0.0

    # ---- チャート用データ(X列〜) ----
    hdr_font = Font(name="ＭＳ Ｐゴシック", size=9, bold=True)
    small = Font(name="ＭＳ Ｐゴシック", size=9)
    c = DATA_COL
    cols = {"km": c}
    heads = ["キロ程"]
    c += 1
    for it in ITEMS:
        cols[f"{it}前回"] = c; heads.append(f"{it}前回"); c += 1
        cols[f"{it}今回"] = c; heads.append(f"{it}今回"); c += 1
    if with_shift:
        cols["km補正"] = c; heads.append(f"前回キロ程(補正{d:+.1f}m)"); c += 1
    # 抽出箇所の縦線: x列(低/高) + 各項目のy列
    cols["vxL"] = c; heads.append("縦線x(10-15mm)"); c += 1
    cols["vxH"] = c; heads.append("縦線x(15mm以上)"); c += 1
    for it in ITEMS:
        cols[f"vy{it}"] = c; heads.append(f"縦線y {it}"); c += 1

    HR = 2  # 見出し行
    ws.cell(row=HR - 1 if HR > 1 else 1, column=DATA_COL)
    for i, hname in enumerate(heads):
        cell = ws.cell(row=HR, column=DATA_COL + i, value=hname)
        cell.font = hdr_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(DATA_COL + i)].width = 9.5
    r = HR
    for k in range(lo, hi + 1):
        r += 1
        v = data.get(k)
        ws.cell(row=r, column=cols["km"], value=k).number_format = KM_FMT
        if with_shift:
            ws.cell(row=r, column=cols["km補正"], value=round(k + d, 1)).number_format = '0.0'
        if v:
            for it in ITEMS:
                ws.cell(row=r, column=cols[f"{it}前回"], value=v[f"{it}前回"])
                ws.cell(row=r, column=cols[f"{it}今回"], value=v[f"{it}今回"])
    r_last = r

    # y軸の範囲(項目ごと、±20mmを基本に、超える場合は10mm単位で拡大)
    yr = {}
    for it in ITEMS:
        m = 0.0
        for k in range(lo, hi + 1):
            v = data.get(k)
            if v:
                for key in (f"{it}前回", f"{it}今回"):
                    if v[key] is not None:
                        m = max(m, abs(v[key]))
        yr[it] = max(20, int(math.ceil(m / 10.0)) * 10)

    # 縦線データ(点(k,-R),(k,+R),空白 の繰り返し)
    vr = r_last + 2
    ws.cell(row=vr - 1, column=cols["vxL"], value="↓縦線用").font = hdr_font
    v0 = vr
    lows = [p["km"] for p in g["points"] if p["判定"] != "15mm以上"]
    highs = [p["km"] for p in g["points"] if p["判定"] == "15mm以上"]
    n = max(len(lows), len(highs))
    for i in range(n):
        for j, sign in enumerate((-1, 1)):
            rr = vr + i * 3 + j
            if i < len(lows):
                ws.cell(row=rr, column=cols["vxL"], value=lows[i]).number_format = KM_FMT
            if i < len(highs):
                ws.cell(row=rr, column=cols["vxH"], value=highs[i]).number_format = KM_FMT
            for it in ITEMS:
                ws.cell(row=rr, column=cols[f"vy{it}"], value=sign * yr[it])
    v1 = vr + n * 3 - 2
    for row in ws.iter_rows(min_row=HR + 1, max_row=v1, min_col=DATA_COL, max_col=c - 1):
        for cell in row:
            cell.font = small

    # ---- チャート(7段) ----
    row0 = 1  # 0始まりの行番号(2行目)
    for idx, it in enumerate(ITEMS):
        last = idx == len(ITEMS) - 1
        ch = ScatterChart()
        ch.style = None
        ch.scatterStyle = "lineMarker"
        ch.display_blanks = "gap"
        ch.series.append(line_series(ws, cols["km"], cols[f"{it}前回"], HR + 1, r_last,
                                     "前回", COLOR_PREV, 0.75))
        if with_shift:
            ch.series.append(line_series(ws, cols["km補正"], cols[f"{it}前回"], HR + 1, r_last,
                                         f"前回(位置補正{d:+.1f}m)", COLOR_SHIFT, 1.0, "dash"))
        ch.series.append(line_series(ws, cols["km"], cols[f"{it}今回"], HR + 1, r_last,
                                     "今回", COLOR_CURR, 2.0))
        if lows:
            ch.series.append(line_series(ws, cols["vxL"], cols[f"vy{it}"], v0, v1,
                                         "抽出(10mm以上15mm未満)", COLOR_LOW, 1.5))
        if highs:
            ch.series.append(line_series(ws, cols["vxH"], cols[f"vy{it}"], v0, v1,
                                         "抽出(15mm以上)", COLOR_HIGH, 1.5))

        x = ch.x_axis
        x.scaling.min, x.scaling.max = lo, hi
        x.majorUnit = 10
        x.number_format = KM_FMT
        x.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="D9D9D9")))
        x.delete = False
        x.axPos = "b"
        x.tickLblPos = "low"
        if not last:
            x.number_format = ";;;"   # 最下段以外は目盛ラベルを非表示(目盛線は表示)
        x.txPr = rich(700)
        x.title = None
        y = ch.y_axis
        y.scaling.min, y.scaling.max = -yr[it], yr[it]
        y.majorUnit = 10 if yr[it] <= 20 else 20
        y.number_format = "0"
        y.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="D9D9D9")))
        y.delete = False
        y.tickLblPos = "low"
        y.txPr = rich(700)
        y.title = title(f"{it}\n[mm]", 800, True)

        # グラフ上部の文字(最上段は線路名と凡例、該当項目は抽出箇所の進み量)
        texts = []
        if idx == 0:
            head = f"{g['track'][0]} {g['track'][1]}({g['track'][2]})　前回:赤細線　今回:青太線"
            if with_shift:
                head += "　前回(位置補正後):緑破線　｜" + shift_info["summary"]
            texts.append(head)
        key = "item_text_shift" if with_shift else "item_text"
        if it in shift_info[key]:
            texts.append(shift_info[key][it])
        if texts:
            ch.title = title("\n".join(texts), 800, True, "006400" if with_shift else None)
            ch.title.overlay = True
        else:
            ch.title = None
            ch.auto_title_deleted = True

        h_in = 0.82 if not last else 0.52
        ch.plot_area.layout = Layout(manualLayout=ManualLayout(
            layoutTarget="inner", xMode="edge", yMode="edge",
            x=0.065, y=0.14 if not last else 0.08, w=0.925, h=h_in))
        if last:
            ch.legend = Legend()
            ch.legend.position = "b"
            ch.legend.txPr = rich(700)
        else:
            ch.legend = None
        ch.anchor = anchor(0, row0, 22, row0 + PANEL_ROWS[idx])
        ws.add_chart(ch)
        row0 += PANEL_ROWS[idx]
    return ws


def write_index(wb, groups, data_by_track, with_shift, shift_by_group, meta):
    ws = wb.active
    ws.title = "一覧"
    bold = Font(name="ＭＳ Ｐゴシック", size=11, bold=True)
    norm = Font(name="ＭＳ Ｐゴシック", size=11)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = f"軌道変位進みの有無確認チャート({'位置補正あり' if with_shift else '位置補正なし'})"
    ws["A1"].font = Font(name="ＭＳ Ｐゴシック", size=14, bold=True)
    ws["A2"] = (f"前回 実績番号 {meta['前回']} ／ 今回 実績番号 {meta['今回']}　"
                f"抽出箇所の前後{meta['window']}mを表示。赤細線=前回、青太線=今回、"
                "縦線=抽出キロ程(オレンジ:10mm以上15mm未満、赤:15mm以上)")
    ws["A2"].font = norm
    r = 3
    if with_shift:
        ws.cell(row=r, column=1, value=(
            f"位置補正:前回の波形を±{meta['max_shift']}mの範囲で0.1mずつずらし、"
            "抽出箇所の前後30mで高低・通り(左右)が今回と最も一致する位置を自動で探した結果(参考値)。"
            f"ズレ量{SHIFT_MIN_M}m以上かつ一致度{SHIFT_MIN_IMPROVE:.0%}以上改善で「位置ズレの可能性あり」。"
            "緑破線=前回(位置補正後)。最終判断は必ずチャートを目視して行うこと。")).font = norm
        r += 1
    r += 1
    heads = ["番号", "シート", "線名", "駅名称", "線路名", "本準側", "キロ程", "判定", "該当項目", "進み量(mm)"]
    if with_shift:
        heads += ["推定位置ズレ(m)", "一致度の改善", "位置ズレの判定(参考)",
                  "補正後の進み量(mm)", "補正後の判定"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(row=r, column=i, value=h)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fill_low = PatternFill("solid", fgColor="FFC000")
    fill_high = PatternFill("solid", fgColor="FF6B6B")
    for gi, g in enumerate(groups):
        sname = sheet_name([p["番号"] for p in g["points"]])
        for p in g["points"]:
            r += 1
            adv_txt = "、".join(f"{it} {v:+.1f}" for it, v in p["差分"].items())
            vals = [p["番号"], sname, p["track"][0], p["駅名称"], p["track"][1], p["track"][2],
                    p["km"], p["判定"], "、".join(p["該当項目"]), adv_txt]
            if with_shift:
                si = shift_by_group[gi]
                ca = si["adv"][p["番号"]]
                worst = max((abs(v) for v in ca.values()), default=0)
                vals += [si["d"], si["improve"], si["verdict"],
                         "、".join(f"{it} {v:+.1f}" for it, v in ca.items()), judge(worst)]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.font = norm
                cell.border = box
            ws.cell(row=r, column=7).number_format = KM_FMT
            ws.cell(row=r, column=8).fill = fill_high if p["判定"] == "15mm以上" else fill_low
            ws.cell(row=r, column=2).hyperlink = f"#'{sname}'!A1"
            ws.cell(row=r, column=2).font = Font(name="ＭＳ Ｐゴシック", size=11, color="0563C1", underline="single")
            if with_shift:
                ws.cell(row=r, column=12).number_format = '0%'
                jc = ws.cell(row=r, column=15)
                if jc.value == "15mm以上":
                    jc.fill = fill_high
                elif jc.value == "10mm以上15mm未満":
                    jc.fill = fill_low
                ws.cell(row=r, column=11).number_format = '+0.0;-0.0;0.0'
    widths = [6, 8, 10, 9, 9, 8, 10, 19, 26, 32, 10, 9, 22, 32, 19]
    for i, w in enumerate(widths[:len(heads)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=6 if with_shift else 5, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def build(src, outdir, window, max_shift):
    tracks, points = load(src)
    if not points:
        raise SystemExit("抽出_要対応 に行がありません")
    groups = make_groups(points, window)
    meta = {"前回": points[0]["前回実績番号"], "今回": points[0]["今回実績番号"],
            "window": window, "max_shift": max_shift}

    # 位置ズレ推定(グループ単位)
    shift_by_group = []
    for g in groups:
        data = tracks[g["track"]]
        d, rms0, rms1 = estimate_shift(data, g["points"], max_shift)
        adv = {p["番号"]: corrected_adv(data, p, d) for p in g["points"]}
        improve = (1 - rms1 / rms0) if rms0 > 0 else 0.0
        shifted = abs(d) >= SHIFT_MIN_M and improve >= SHIFT_MIN_IMPROVE
        verdict = "位置ズレの可能性あり" if shifted else "位置ズレは見られない"
        summary = (f"推定位置ズレ {d:+.1f}m(一致度 {improve:.0%} 改善)→{verdict}")
        item_text, item_text_shift = {}, {}
        for it in ITEMS:
            before = [abs(p["差分"][it]) for p in g["points"] if it in p["差分"]]
            after = [abs(adv[p["番号"]][it]) for p in g["points"] if it in adv[p["番号"]]]
            if before:
                item_text[it] = f"{it}:抽出箇所の進み量 最大{max(before):.1f}mm"
            if before and after:
                item_text_shift[it] = (f"{it}:抽出箇所の進み量 最大{max(before):.1f}mm → "
                                       f"位置補正({d:+.1f}m)後 最大{max(after):.1f}mm")
        shift_by_group.append({"d": d, "rms0": rms0, "rms1": rms1, "adv": adv,
                               "improve": improve, "verdict": verdict, "summary": summary,
                               "item_text": item_text, "item_text_shift": item_text_shift})
        print(f"group {sheet_name([p['番号'] for p in g['points']])}: shift {d:+.1f}m "
              f"RMS {rms0:.2f} -> {rms1:.2f}")

    os.makedirs(outdir, exist_ok=True)
    outs = []
    for with_shift in (False, True):
        wb = openpyxl.Workbook()
        write_index(wb, groups, tracks, with_shift, shift_by_group, meta)
        for gi, g in enumerate(groups):
            write_group_sheet(wb, g, tracks[g["track"]], window, with_shift, shift_by_group[gi])
        label = "位置補正あり" if with_shift else "位置補正なし"
        fn = os.path.join(outdir, f"【{meta['今回']}】軌道変位進みチャート_{label}.xlsx")
        wb.save(fn)
        outs.append(fn)
        print("saved:", fn)
    return outs, groups, shift_by_group


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("outdir")
    ap.add_argument("--window", type=int, default=100)
    ap.add_argument("--max-shift", type=float, default=5.0)
    a = ap.parse_args()
    build(a.src, a.outdir, a.window, a.max_shift)
